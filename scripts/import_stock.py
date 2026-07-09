"""Импорт остатков склада из Excel (инвентаризация) в МойСклад.

Приводит физический остаток склада к колонке «Текущий остаток» листа
«Остатки»: недостающее оприходуется одним документом (себестоимость —
закупочная из каталога Uzum). Позиции, где в МойСклад больше, чем в
файле, НЕ списываются — только отчёт (списание запускается осознанно).

Запуск:
    python scripts/import_stock.py <файл.xlsx> [--store source|uzum] [--dry-run]
"""
from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from config import PRICE_MINOR_UNITS  # noqa: E402
from main import build_context  # noqa: E402
from moysklad.entities import meta_ref  # noqa: E402

log = logging.getLogger("import_stock")


def read_sheet(path: str) -> dict[str, int]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Остатки"]
    header = [c.value for c in ws[1]]
    sku_col = header.index("SKU")
    qty_col = header.index("Текущий остаток")
    out: dict[str, int] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        sku = row[sku_col]
        qty = row[qty_col]
        if not sku or qty is None:
            continue
        sku = str(sku).strip()
        if sku in out:
            # дубли — обычно «хвосты» без начального остатка; основная строка первая
            log.warning("Дубль SKU «%s» в файле: беру первую строку (%s), игнорирую %s",
                        sku, out[sku], int(qty))
            continue
        out[sku] = int(qty)
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx")
    parser.add_argument("--store", choices=["source", "uzum"], default="source")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)-7s %(message)s")
    logging.getLogger("httpx").setLevel(logging.WARNING)

    ctx = build_context()
    store = ctx.entities.store_source() if args.store == "source" else ctx.entities.store_uzum()
    store_id = store["meta"]["href"].rstrip("/").rsplit("/", 1)[-1]
    log.info("Склад: «%s»", store.get("name"))

    sheet = read_sheet(args.xlsx)
    log.info("В файле %d SKU (ненулевых: %d)", len(sheet), sum(1 for q in sheet.values() if q))

    # каталог Uzum: skuTitle -> CatalogSku (для штрихкода и закупочной)
    by_title: dict[str, list] = {}
    for sku in ctx.resolver.catalog_skus():
        by_title.setdefault(sku.sku_title, []).append(sku)

    rows = ctx.ms.get("/report/stock/bystore/current", {"stockType": "stock"}) or []
    ms_stock = {r["assortmentId"]: r["stock"] for r in rows if r.get("storeId") == store_id}

    positions, to_write_off, unmatched = [], [], []
    for title, target in sheet.items():
        candidates = by_title.get(title) or []
        found = None
        for c in candidates:
            found = ctx.resolver.resolve(c.sku_title, product_id=c.product_id, sku_id=c.sku_id)
            if found:
                break
        if found is None:
            found = ctx.entities.find_assortment(title)
        if found is None:
            if target:
                unmatched.append((title, target))
            continue
        aid = found["href"].rstrip("/").rsplit("/", 1)[-1]
        current = int(ms_stock.get(aid, 0))
        delta = target - current
        if delta > 0:
            purchase = max((c.purchase_price for c in candidates), default=0)
            pos = {
                "assortment": {"meta": {"href": found["href"], "type": found["type"],
                                        "mediaType": "application/json"}},
                "quantity": delta,
            }
            if purchase:
                pos["price"] = purchase * PRICE_MINOR_UNITS
            positions.append(pos)
            log.info("«%s»: файл %d, МойСклад %d -> оприходовать %+d", title, target, current, delta)
        elif delta < 0:
            to_write_off.append((title, target, current))

    for title, target in unmatched:
        log.warning("НЕ НАЙДЕН в МойСклад: «%s» (остаток по файлу %d)", title, target)
    for title, target, current in to_write_off:
        log.warning("В МойСклад БОЛЬШЕ файла: «%s» файл %d, МС %d — не списываю", title, target, current)

    if not positions:
        log.info("Оприходовать нечего")
        return 0
    total = sum(p["quantity"] for p in positions)
    if args.dry_run:
        log.info("[dry-run] Оприходование: %d позиций, %d шт", len(positions), total)
        return 0

    doc = ctx.ms.post(
        "/entity/enter",
        {
            "organization": meta_ref(ctx.entities.organization()),
            "store": meta_ref(store),
            "description": f"Оприходование остатков склада «{store.get('name')}» "
                           f"по файлу инвентаризации ({Path(args.xlsx).name})",
            "positions": positions,
        },
    )
    log.info("Создано оприходование %s: %d позиций, %d шт", doc.get("name"), len(positions), total)
    return 0


if __name__ == "__main__":
    sys.exit(main())
